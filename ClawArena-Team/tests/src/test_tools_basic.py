from __future__ import annotations

import os
import shutil
import time
from pathlib import Path

import pytest

from clawarena_team.sandbox import AccessibleScope, ReadTracker
from clawarena_team.tools.base import ToolContext
from clawarena_team.tools.basic import BashTool, EditTool, GlobTool, GrepTool, ReadTool, WriteTool


def _ctx(tmp_workspace: Path, *, modalities=None) -> ToolContext:
    scope = AccessibleScope(
        [tmp_workspace / "inbox", tmp_workspace / "notes"], scenario_root=tmp_workspace
    )
    return ToolContext(
        agent_id="test",
        scope=scope,
        read_tracker=ReadTracker(),
        cwd=tmp_workspace / "inbox",
        config={
            "read.max_text_bytes": 524288,
            "bash.default_timeout_ms": 10000,
            "bash.max_output_bytes": 65536,
        },
        agent_modalities=modalities or ["text"],
    )


async def test_read_text(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run(
        {"file_path": str(tmp_workspace / "inbox" / "hello.txt")}, ctx
    )
    assert not res.is_error
    assert "hello world" in res.content
    assert "Read" in ctx.tools_used


async def test_read_requires_absolute(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run({"file_path": "hello.txt"}, ctx)
    assert res.is_error
    assert "absolute" in res.content.lower()


async def test_read_forbidden(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run(
        {"file_path": str(tmp_workspace / "restricted" / "secret.txt")}, ctx
    )
    assert res.is_error
    assert "forbidden" in res.content


async def test_read_image_native_vs_placeholder(tmp_workspace: Path):
    img = tmp_workspace / "inbox" / "x.png"
    img.write_bytes(b"\x89PNG\r\n")
    ctx_text = _ctx(tmp_workspace, modalities=["text"])
    res = await ReadTool().run({"file_path": str(img)}, ctx_text)
    assert "not natively supported" in res.content
    # 错误信息必须按 modality 提示而非 model_key
    assert "`image` modality" in res.content
    assert "model_key" not in res.content
    ctx_vlm = _ctx(tmp_workspace, modalities=["text", "image"])
    res2 = await ReadTool().run({"file_path": str(img)}, ctx_vlm)
    assert "natively supported" in res2.content
    assert any(a[1] == "image" for a in res2.attachments)


def test_read_schema_extension_listing_per_modality():
    schema_text = ReadTool.schema(modalities=["text"])["description"]
    # 始终列出 office 与四类文本分组
    assert "office (auto-parsed" in schema_text
    assert "data / config" in schema_text and ".csv" in schema_text
    assert "code" in schema_text and ".go" in schema_text and ".sql" in schema_text
    # 主代理 modality 为 text 时，image/audio/video 都在"Not natively supported"块中
    assert "Not natively supported" in schema_text
    assert "image:" in schema_text and "audio:" in schema_text and "video:" in schema_text

    schema_vlm = ReadTool.schema(modalities=["text", "image"])["description"]
    # 已声明 image 时不应再列入 unsupported 块
    assert "image (native)" in schema_vlm
    after_native = schema_vlm.split("Not natively supported", 1)
    if len(after_native) == 2:
        assert "image:" not in after_native[1].split("- A successful Read")[0]
    # model_key 不应出现在 Read schema 文案里（用户要求 model_key 仅供 subagent 工具）
    assert "model_key" not in schema_vlm


async def test_read_office_xlsx(tmp_workspace: Path):
    import openpyxl

    xlsx = tmp_workspace / "inbox" / "report.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["region", "q1", "q2"])
    ws.append(["west", 1200, 1450])
    wb.save(xlsx)
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run({"file_path": str(xlsx)}, ctx)
    assert not res.is_error
    assert "Sales" in res.content and "1200" in res.content
    assert any(a[1] == "text" for a in res.attachments)


async def test_read_office_docx(tmp_workspace: Path):
    from docx import Document

    docx = tmp_workspace / "inbox" / "memo.docx"
    doc = Document()
    doc.add_paragraph("First line.")
    doc.add_paragraph("Second line: 42 widgets.")
    doc.save(docx)
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run({"file_path": str(docx)}, ctx)
    assert not res.is_error
    assert "First line." in res.content and "42 widgets" in res.content


async def test_read_office_pdf(tmp_workspace: Path):
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas as pdfcanvas

    pdf = tmp_workspace / "inbox" / "doc.pdf"
    c = pdfcanvas.Canvas(str(pdf))
    c.drawString(72, 720, "Quarter revenue 1234.")
    c.save()
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run({"file_path": str(pdf)}, ctx)
    assert not res.is_error
    assert "Quarter revenue 1234" in res.content


async def test_read_unknown_extension_treated_as_text(tmp_workspace: Path):
    # 后缀未被任何集合命中也应按文本读取（"everything else" 兜底）
    p = tmp_workspace / "inbox" / "weird.xyz"
    p.write_text("plain text payload here", encoding="utf-8")
    ctx = _ctx(tmp_workspace)
    res = await ReadTool().run({"file_path": str(p)}, ctx)
    assert not res.is_error
    assert "plain text payload" in res.content


async def test_write_new_file_no_read_required(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await WriteTool().run(
        {"file_path": str(tmp_workspace / "inbox" / "new.md"), "content": "alpha"}, ctx
    )
    assert not res.is_error
    assert (tmp_workspace / "inbox" / "new.md").read_text() == "alpha"


async def test_write_overwrite_requires_read(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await WriteTool().run(
        {"file_path": str(tmp_workspace / "inbox" / "hello.txt"), "content": "x"}, ctx
    )
    assert res.is_error
    assert "read" in res.content.lower()


async def test_edit_requires_prior_read(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await EditTool().run(
        {
            "file_path": str(tmp_workspace / "inbox" / "hello.txt"),
            "old_string": "world",
            "new_string": "earth",
        },
        ctx,
    )
    assert res.is_error


async def test_edit_after_read(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    path = str(tmp_workspace / "inbox" / "hello.txt")
    await ReadTool().run({"file_path": path}, ctx)
    res = await EditTool().run(
        {"file_path": path, "old_string": "world", "new_string": "earth"}, ctx
    )
    assert not res.is_error
    assert (tmp_workspace / "inbox" / "hello.txt").read_text() == "hello earth"


async def test_bash_cwd_and_capture(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await BashTool().run({"command": "pwd", "description": "Show cwd"}, ctx)
    assert not res.is_error
    assert str(tmp_workspace / "inbox") in res.content


async def test_bash_timeout_ms(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await BashTool().run(
        {"command": "sleep 2", "description": "Sleep 2 seconds", "timeout": 100}, ctx
    )
    assert res.is_error
    assert "100ms" in res.content


@pytest.mark.skipif(shutil.which("rg") is None, reason="ripgrep not installed")
async def test_grep_default_files_with_matches(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await GrepTool().run(
        {"pattern": "hello", "path": str(tmp_workspace / "inbox")}, ctx
    )
    assert "hello.txt" in res.content


@pytest.mark.skipif(shutil.which("rg") is None, reason="ripgrep not installed")
async def test_grep_content_mode(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    res = await GrepTool().run(
        {
            "pattern": "hello",
            "path": str(tmp_workspace / "inbox"),
            "output_mode": "content",
            "-n": True,
        },
        ctx,
    )
    assert "hello world" in res.content


async def test_glob_sort_mtime(tmp_workspace: Path):
    ctx = _ctx(tmp_workspace)
    older = tmp_workspace / "inbox" / "older.txt"
    older.write_text("o")
    os.utime(older, (time.time() - 1000, time.time() - 1000))
    newer = tmp_workspace / "inbox" / "newer.txt"
    newer.write_text("n")
    res = await GlobTool().run(
        {"pattern": "*.txt", "path": str(tmp_workspace / "inbox")}, ctx
    )
    lines = res.content.splitlines()
    idx_new = next(i for i, l in enumerate(lines) if l.endswith("newer.txt"))
    idx_old = next(i for i, l in enumerate(lines) if l.endswith("older.txt"))
    assert idx_new < idx_old
