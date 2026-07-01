"""Internal helpers shared by basic tools.

集中维护：

- 扩展名分类与 :func:`detect_modality`
- :func:`parse_office` —— xlsx / docx / pdf 透明解析为纯文本
- :func:`require_absolute` / :func:`forbidden` —— Read/Write/Edit/Grep/Glob 共用的
  路径前置校验
- :func:`params` —— 缩写访问 ``prompts.TOOL_PARAM_DESCRIPTIONS[<tool>]``

仅供 ``arcbench.tools`` 包内部使用；公共导出在 :mod:`arcbench.tools.__init__`
集中。文件名前缀下划线表示 private（``detect_modality`` 通过 ``__init__`` 重导）。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from ..prompts import (
    ABSOLUTE_PATH_REQUIRED_MESSAGE,
    FORBIDDEN_PATH_MESSAGE,
    TOOL_PARAM_DESCRIPTIONS,
)
from .base import ToolExecResult


# ---------------------------------------------------------------------------
# Modality / extension detection
# ---------------------------------------------------------------------------

DOC_EXTS = {".md", ".markdown", ".txt", ".rst", ".rtf"}
DATA_EXTS = {
    ".csv", ".tsv", ".json", ".jsonl", ".ndjson",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env",
    ".xml", ".html", ".htm", ".svg",
}
LOG_EXTS = {".log"}
CODE_EXTS = {
    ".py", ".pyi", ".ipynb",
    ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx",
    ".go", ".rs", ".java", ".kt", ".scala", ".swift",
    ".rb", ".php", ".pl", ".lua",
    ".c", ".cc", ".cpp", ".cxx", ".h", ".hpp", ".cs",
    ".sh", ".bash", ".zsh", ".fish",
    ".sql", ".graphql", ".proto",
    ".tf", ".hcl",
    ".vue", ".svelte",
    ".dockerfile", ".mk", ".makefile",
    ".patch", ".diff",
}
OFFICE_EXTS = {".xlsx", ".docx", ".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff"}
AUDIO_EXTS = {".wav", ".mp3", ".flac", ".ogg", ".m4a", ".aac"}
VIDEO_EXTS = {".mp4", ".mov", ".avi", ".mkv", ".webm"}


def detect_modality(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in IMAGE_EXTS:
        return "image"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in OFFICE_EXTS:
        return "office"
    return "text"


# ---------------------------------------------------------------------------
# Office parsers
# ---------------------------------------------------------------------------


def parse_office(path: Path) -> tuple[str, str | None]:
    ext = path.suffix.lower()
    try:
        if ext == ".xlsx":
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                parts.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    parts.append("\t".join("" if v is None else str(v) for v in row))
                parts.append("")
            wb.close()
            return "\n".join(parts), None
        if ext == ".docx":
            from docx import Document  # type: ignore

            doc = Document(str(path))
            parts = []
            for para in doc.paragraphs:
                if para.text:
                    parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    parts.append("\t".join(cells))
            return "\n".join(parts), None
        if ext == ".pdf":
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(str(path))
            parts = []
            for i, page in enumerate(reader.pages):
                parts.append(f"--- Page {i + 1} ---")
                parts.append(page.extract_text() or "")
            return "\n".join(parts), None
    except ImportError as e:
        return "", (
            f"office parser dependency missing for {ext}: {e.name or e}. "
            f"Install one of: openpyxl (.xlsx) / python-docx (.docx) / pypdf (.pdf)."
        )
    except Exception as e:  # noqa: BLE001
        return "", f"office parse error for {path.name}: {type(e).__name__}: {e}"
    return "", f"unsupported office extension: {ext}"


# ---------------------------------------------------------------------------
# Path validation
# ---------------------------------------------------------------------------


def require_absolute(p: str) -> Path | ToolExecResult:
    pp = Path(p)
    if not pp.is_absolute():
        return ToolExecResult(
            ABSOLUTE_PATH_REQUIRED_MESSAGE.format(path=p), is_error=True
        )
    return pp.resolve(strict=False)


def forbidden(path: str) -> ToolExecResult:
    return ToolExecResult(FORBIDDEN_PATH_MESSAGE.format(path=path), is_error=True)


# ---------------------------------------------------------------------------
# Param description shortcut
# ---------------------------------------------------------------------------


def params(tool: str) -> dict[str, str]:
    return TOOL_PARAM_DESCRIPTIONS[tool]


def schema_format_kwargs(
    config_dict: dict[str, Any], modalities: list[str] | tuple[str, ...]
) -> dict[str, Any]:
    """构造工具 schema 描述模板的填充参数，使描述与运行时配置一致。

    供 main / subagent 两处调用 ``cls.schema(**...)`` 共用，避免描述里写死阈值与
    超时而与 ``read.*`` / ``bash.*`` 配置漂移。各工具 ``schema`` 用 ``.get`` 取需要的
    键，未传则回落默认。
    """
    read = config_dict.get("read", {}) or {}
    bash = config_dict.get("bash", {}) or {}
    return {
        "modalities": list(modalities),
        "notice_bytes": int(read.get("notice_bytes", 32768)),
        "hard_bytes": int(read.get("hard_bytes", 262144)),
        "bash_default_timeout_ms": int(bash.get("default_timeout_sec", 60)) * 1000,
    }


__all__ = [
    "DOC_EXTS",
    "DATA_EXTS",
    "LOG_EXTS",
    "CODE_EXTS",
    "OFFICE_EXTS",
    "IMAGE_EXTS",
    "AUDIO_EXTS",
    "VIDEO_EXTS",
    "detect_modality",
    "parse_office",
    "require_absolute",
    "forbidden",
    "params",
    "schema_format_kwargs",
]
